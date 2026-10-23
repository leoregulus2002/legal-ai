import re
from collections.abc import Sequence
from dataclasses import dataclass, replace
from decimal import Decimal
from io import BytesIO
from itertools import pairwise

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.slide import Slide as PptxSlide
from pptx.text.text import Font, TextFrame
from pptx.util import Pt
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

from app.common.exceptions import AppError
from app.modules.files.schemas import FileSpec, Slide, SlideChart, SlideColumn, SlideStep

WIDTH, HEIGHT, MARGIN = 960, 540, 52
BODY_WIDTH = WIDTH - MARGIN * 2
FONT = "Microsoft YaHei"
INK, PAPER, GOLD, MUTED, WHITE = "172938", "F8F6F0", "987545", "667582", "FFFFFF"
PALETTE = (INK, GOLD, "547D89", "9EADB4")
MAX_PAGES = 240
# 使用跨平台可用的中文字体度量，并预留无衬线字体差异的空间；不依赖服务器安装 Office 字体。
_metrics = UnicodeCIDFont("STSong-Light")


@dataclass(frozen=True)
class Page:
    entry: Slide
    layout: str
    part: int = 1
    parts: int = 1
    step_numbers: tuple[int, ...] = ()
    column_widths: tuple[float, ...] = ()


def _width(text: str, size: float) -> float:
    return _metrics.stringWidth(text.expandtabs(4), size) * 1.12


def _wrap(text: str, width: float, size: float) -> list[str]:
    lines: list[str] = []
    normalized = text.expandtabs(4).replace("\r\n", "\n").replace("\r", "\n").replace("\v", "\n")
    for paragraph in normalized.split("\n"):
        line = ""
        for token in re.findall(r"[A-Za-z0-9][A-Za-z0-9_./:@%+\-]*|\s+|.", paragraph):
            pieces = list(token) if _width(token, size) > width else [token]
            for piece in pieces:
                if line and _width(line + piece, size) > width:
                    lines.append(line.rstrip())
                    line = ""
                line += piece
        lines.append(line.rstrip())
    return lines or [""]


def _height(text: str, width: float, size: float) -> float:
    return len(_wrap(text, width, size)) * size * 1.32 + 3


def _font(font: Font, size: float, color: str = INK, *, bold: bool = False) -> None:
    font.name, font.size, font.bold = FONT, Pt(size), bold
    font.color.rgb = RGBColor.from_string(color)
    # python-pptx 的 name 只设置西文字体；中文必须同时指定东亚字体。
    east_asian = font._rPr.find(qn("a:ea"))
    if east_asian is None:
        east_asian = OxmlElement("a:ea")
        font._rPr.append(east_asian)
    east_asian.set("typeface", FONT)


def _frame(
    frame: TextFrame,
    text: str,
    width: float,
    size: float,
    color: str,
    *,
    bold: bool = False,
    bullet: bool = False,
) -> None:
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap, frame.auto_size = False, MSO_AUTO_SIZE.NONE
    paragraph = frame.paragraphs[0]
    paragraph.text = "\v".join(_wrap(text, width - (22 if bullet else 0), size))
    paragraph.line_spacing = Pt(size * 1.32)
    paragraph.space_before = paragraph.space_after = Pt(0)
    _font(paragraph.font, size, color, bold=bold)
    if bullet:
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Pt(22)))
        properties.set("indent", str(-Pt(16)))
        marker = OxmlElement("a:buChar")
        marker.set("char", "•")
        properties.append(marker)


def _text(
    slide: PptxSlide,
    text: str,
    x: float,
    y: float,
    width: float,
    size: float,
    color: str = INK,
    *,
    bold: bool = False,
    bullet: bool = False,
) -> float:
    height = _height(text, width - (22 if bullet else 0), size)
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
    _frame(box.text_frame, text, width, size, color, bold=bold, bullet=bullet)
    return height


def _line(
    slide: PptxSlide, x1: float, y1: float, x2: float, y2: float, *, arrow: bool = False
) -> None:
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    connector.line.color.rgb, connector.line.width = RGBColor.from_string(GOLD), Pt(1.3)
    if arrow:
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        connector.line._get_or_add_ln().append(tail)


def _header(entry: Slide) -> tuple[float, float, float]:
    size = 34 if len(_wrap(entry.title, BODY_WIDTH, 34)) <= 3 else 32
    top = 42 + _height(entry.title, BODY_WIDTH, size)
    if entry.subtitle:
        top += 12 + _height(entry.subtitle, BODY_WIDTH, 18)
    top += 26
    bottom = 466 if entry.source else 488
    if bottom - top < 65:
        raise AppError("presentation_heading_too_long", "幻灯片标题或副标题过长，请精简后重试")
    return top, bottom, size


def _cover(entry: Slide) -> tuple[float, float, float]:
    size = next(
        (size for size in (44, 38, 32) if len(_wrap(entry.title, BODY_WIDTH, size)) <= 4), 32
    )
    total = _height(entry.title, BODY_WIDTH, size)
    if entry.subtitle:
        total += 22 + _height(entry.subtitle, BODY_WIDTH, 22)
    top = max(64, min(145, (440 - total) / 2))
    if top + total > 450:
        raise AppError("presentation_heading_too_long", "幻灯片标题或副标题过长，请精简后重试")
    return top, top + total + 30, size


def _pieces(text: str, width: float, size: float, height: float) -> list[str]:
    if height < size * 1.32 + 3:
        raise AppError("presentation_content_too_dense", "幻灯片内容过密，请减少单页内容后重试")
    result: list[str] = []
    current = ""
    # 优先在完整句子之间分页，只有单句超长时才按可用行数拆分，不丢弃正文。
    for sentence in re.split(r"(?<=[。！？；])", text):
        if _height(current + sentence, width, size) <= height:
            current += sentence
            continue
        if current:
            result.append(current)
            current = ""
        if _height(sentence, width, size) <= height:
            current = sentence
            continue
        lines = _wrap(sentence, width, size)
        count = max(1, int((height - 3) / (size * 1.32)))
        result.extend(
            "\n".join(lines[index : index + count]) for index in range(0, len(lines), count)
        )
    if current:
        result.append(current)
    return result or [text]


def _pack(
    values: list[str], width: float, size: float, height: float, *, gap: float = 16
) -> list[list[str]]:
    pages: list[list[str]] = [[]]
    used = 0.0
    for value in values:
        if not value.strip():
            continue
        for piece in _pieces(value, width, size, height):
            needed = _height(piece, width, size)
            if pages[-1] and used + gap + needed > height:
                pages.append([])
                used = 0
            used += needed + (gap if pages[-1] else 0)
            pages[-1].append(piece)
    return pages


def _layout(entry: Slide) -> str:
    if entry.layout != "auto":
        return entry.layout
    if entry.columns:
        return "comparison"
    if entry.steps:
        return "process"
    if entry.quote:
        return "quote"
    if entry.rows:
        return "table"
    if entry.chart.series:
        return "chart"
    return "bullets"


def _column_heading(entry: Slide) -> float:
    return max(_height(column.title, (BODY_WIDTH - 52) / 2, 24) for column in entry.columns) + 22


def _table_widths(rows: list[list[str]]) -> list[float]:
    weights = [
        max(1.0, min(3.0, max(_width(row[column], 18) for row in rows) / 150))
        for column in range(len(rows[0]))
    ]
    return [BODY_WIDTH * weight / sum(weights) for weight in weights]


def _row_height(row: list[str], widths: Sequence[float]) -> float:
    return (
        max(_height(value, width - 20, 18) for value, width in zip(row, widths, strict=True)) + 18
    )


def _table_pages(entry: Slide, height: float) -> list[Slide]:
    widths = _table_widths(entry.rows)
    header, *rows = entry.rows
    header_height = _row_height(header, widths)
    available = height - header_height
    if available < 18 * 1.32 + 21:
        raise AppError("presentation_table_too_dense", "幻灯片表头过长，请精简表头后重试")
    pages: list[list[list[str]]] = [[header]]
    used = header_height
    for row in rows:
        wrapped = [_wrap(value, width - 20, 18) for value, width in zip(row, widths, strict=True)]
        capacity = max(1, int((available - 21) / (18 * 1.32)))
        for start in range(0, max(map(len, wrapped)), capacity):
            values = ["\n".join(lines[start : start + capacity]) for lines in wrapped]
            row_height = _row_height(values, widths)
            if len(pages[-1]) > 1 and used + row_height > height:
                pages.append([header])
                used = header_height
            pages[-1].append(values)
            used += row_height
    return [entry.model_copy(update={"rows": rows}) for rows in pages]


def _horizontal(entry: Slide, height: float) -> bool:
    if not 2 <= len(entry.steps) <= 4:
        return False
    width = (BODY_WIDTH - 44 * (len(entry.steps) - 1)) / len(entry.steps)
    return all(
        _height(step.label, width, 22) + _height(step.text, width, 20) + 85 <= height
        for step in entry.steps
    )


def _step_pages(entry: Slide, height: float) -> list[Page]:
    if entry.layout == "process" and _horizontal(entry, height):
        return [Page(entry, "process")]
    pages: list[list[SlideStep]] = [[]]
    used = 0.0
    numbers: list[list[int]] = [[]]
    for index, step in enumerate(entry.steps):
        label_height = _height(step.label, 170, 20)
        if label_height + 20 > height:
            raise AppError("presentation_step_too_dense", "幻灯片步骤标题过长，请精简后重试")
        for text in _pieces(step.text, BODY_WIDTH - 245, 20, height - 20):
            row_height = max(label_height, _height(text, BODY_WIDTH - 245, 20)) + 20
            if pages[-1] and used + row_height > height:
                pages.append([])
                numbers.append([])
                used = 0
            pages[-1].append(step.model_copy(update={"text": text}))
            numbers[-1].append(index + 1)
            used += row_height
    return [
        Page(entry.model_copy(update={"steps": steps}), _layout(entry), step_numbers=tuple(order))
        for steps, order in zip(pages, numbers, strict=True)
    ]


def _chart_table(entry: Slide) -> Slide:
    chart = entry.chart
    if not chart.series:
        return entry
    unit = f"（{chart.unit}）" if chart.unit else ""
    rows = [["项目", *[series.name + unit for series in chart.series]]]
    for index, category in enumerate(chart.categories):
        values = [series.values[index] for series in chart.series]
        rows.append(
            [
                category,
                *[
                    format((Decimal(str(value)) * 100).normalize(), "f") + "%"
                    if chart.value_format == "percent"
                    else str(value)
                    for value in values
                ],
            ]
        )
    return entry.model_copy(update={"layout": "table", "chart": SlideChart(), "rows": rows})


def _chart_fits(entry: Slide, height: float) -> bool:
    data = entry.chart
    plot_height = height - 65 - (30 if data.unit else 0) - (40 if len(data.series) > 1 else 0)
    if plot_height < 180 or any(_width(label, 16) > 180 for label in data.categories):
        return False
    slots = len(data.categories) * len(data.series)
    if data.kind == "bar":
        return slots * 26 <= plot_height
    value_width = max(
        _width(format(Decimal(str(value)).normalize(), "f"), 14)
        for series in data.series
        for value in series.values
    )
    return (
        slots <= 18
        and slots * max(34, value_width + 16) <= BODY_WIDTH - 140
        and all(
            _height(label, (BODY_WIDTH - 140) / len(data.categories), 14) <= 60
            for label in data.categories
        )
    )


def _pages(entry: Slide) -> list[Page]:
    layout = _layout(entry)
    if layout in {"cover", "section"}:
        _, top, _ = _cover(entry)
        # 封面放不下的正文进入正常内容页，完整保留文字。
        if not entry.points:
            return [Page(entry, layout)]
        if top + sum(_height(point, BODY_WIDTH, 22) + 16 for point in entry.points) <= 466:
            return [Page(entry, layout)]
        return [
            Page(entry.model_copy(update={"points": []}), layout),
            *_pages(entry.model_copy(update={"layout": "bullets"})),
        ]
    top, bottom, _ = _header(entry)
    height = bottom - top
    if layout == "comparison":
        height -= _column_heading(entry)
        if height < 22 * 1.32 + 3:
            raise AppError("presentation_heading_too_long", "对比页标题过长，请精简后重试")
        columns = [
            _pack(column.points, (BODY_WIDTH - 52) / 2 - 22, 22, height) for column in entry.columns
        ]
        result = [
            Page(
                entry.model_copy(
                    update={
                        "columns": [
                            SlideColumn(
                                title=column.title,
                                points=chunks[index] if index < len(chunks) else [],
                            )
                            for column, chunks in zip(entry.columns, columns, strict=True)
                        ]
                    }
                ),
                layout,
            )
            for index in range(max(map(len, columns)))
        ]
    elif layout in {"timeline", "process"}:
        result = _step_pages(entry, height)
    elif layout == "quote":
        reserve = _height(entry.attribution, BODY_WIDTH - 40, 18) + 20 if entry.attribution else 0
        result = [
            Page(entry.model_copy(update={"quote": text}), layout)
            for text in _pieces(entry.quote, BODY_WIDTH - 40, 26, height - reserve)
        ]
    elif layout == "table":
        widths = tuple(_table_widths(entry.rows))
        result = [Page(item, layout, column_widths=widths) for item in _table_pages(entry, height)]
    elif layout == "chart":
        if not _chart_fits(entry, height):
            return _pages(_chart_table(entry))
        result = [Page(entry, layout)]
    else:
        result = [
            Page(entry.model_copy(update={"points": points}), layout)
            for points in _pack(entry.points, BODY_WIDTH - 22, 24, height)
        ]
    return [replace(page, part=index + 1, parts=len(result)) for index, page in enumerate(result)]


def _draw_table(slide: PptxSlide, entry: Slide, top: float, widths: tuple[float, ...]) -> None:
    heights = [_row_height(row, widths) for row in entry.rows]
    table = slide.shapes.add_table(
        len(entry.rows), len(widths), Pt(MARGIN), Pt(top), Pt(BODY_WIDTH), Pt(sum(heights))
    ).table
    table.first_row, table.horz_banding = True, False
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Pt(width)
    for index, values in enumerate(entry.rows):
        table.rows[index].height = Pt(heights[index])
        for column, value in enumerate(values):
            cell = table.cell(index, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(INK if index == 0 else WHITE)
            _frame(
                cell.text_frame,
                value,
                widths[column] - 20,
                18,
                WHITE if index == 0 else INK,
                bold=index == 0,
            )
            cell.margin_left = cell.margin_right = Pt(10)
            cell.margin_top = cell.margin_bottom = Pt(9)


def _draw_steps(slide: PptxSlide, page: Page, top: float, height: float) -> None:
    entry = page.entry
    if page.layout == "process" and _horizontal(entry, height):
        width = (BODY_WIDTH - 44 * (len(entry.steps) - 1)) / len(entry.steps)
        for index, step in enumerate(entry.steps):
            x = MARGIN + index * (width + 44)
            number = page.step_numbers[index] if page.step_numbers else index + 1
            _text(slide, f"{number:02d}", x, top, width, 32, GOLD, bold=True)
            label_height = _text(slide, step.label, x, top + 62, width, 22, bold=True)
            _text(slide, step.text, x, top + 78 + label_height, width, 20)
            if index < len(entry.steps) - 1:
                _line(slide, x + width + 6, top + 22, x + width + 34, top + 22, arrow=True)
        return
    y = top
    centers: list[float] = []
    for index, step in enumerate(entry.steps):
        label_height = _height(step.label, 170, 20)
        text_height = _height(step.text, BODY_WIDTH - 245, 20)
        centers.append(y + 13)
        if page.layout == "timeline":
            node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(MARGIN + 5), Pt(y + 8), Pt(10), Pt(10))
            node.fill.solid()
            node.fill.fore_color.rgb = RGBColor.from_string(GOLD)
            node.line.fill.background()
        else:
            number = page.step_numbers[index] if page.step_numbers else index + 1
            _text(slide, f"{number:02d}", MARGIN, y, 35, 20, GOLD, bold=True)
        _text(slide, step.label, MARGIN + 45, y, 170, 20, bold=True)
        _text(slide, step.text, MARGIN + 245, y, BODY_WIDTH - 245, 20)
        y += max(label_height, text_height) + 20
    if page.layout == "timeline":
        for start, end in pairwise(centers):
            _line(slide, MARGIN + 10, start + 6, MARGIN + 10, end - 6)


def _draw_chart(slide: PptxSlide, entry: Slide, top: float, height: float) -> None:
    data = entry.chart
    if not data.series:
        return
    if data.unit:
        top += _text(slide, f"单位：{data.unit}", MARGIN, top, BODY_WIDTH, 14, MUTED) + 8
        height -= 30
    number_format = "0.##%" if data.value_format == "percent" else "#,##0.##"
    chart_data = CategoryChartData(number_format=number_format)
    chart_data.categories = data.categories
    for series in data.series:
        chart_data.add_series(series.name, series.values)
    chart = slide.shapes.add_chart(
        {
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS,
        }[data.kind],
        Pt(MARGIN),
        Pt(top),
        Pt(BODY_WIDTH),
        Pt(height),
        chart_data,
    ).chart
    chart.has_title, chart.has_legend = False, len(data.series) > 1
    _font(chart.font, 16)
    properties = OxmlElement("c:spPr")
    fill = OxmlElement("a:solidFill")
    fill_color = OxmlElement("a:srgbClr")
    fill_color.set("val", PAPER)
    fill.append(fill_color)
    properties.append(fill)
    chart._chartSpace.insert_element_before(
        properties, "c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst"
    )
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        _font(chart.legend.font, 14)
    for axis in (chart.category_axis, chart.value_axis):
        _font(axis.tick_labels.font, 14, MUTED)
        axis.format.line.fill.background()
    chart.value_axis.tick_labels.number_format = number_format
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("DDDCD6")
    plot = chart.plots[0]
    # 多系列用图例区分，避免相等或接近的数据点标签互相覆盖；原始数值保留在图表工作簿中。
    plot.has_data_labels = len(data.series) == 1
    if plot.has_data_labels:
        plot.data_labels.show_value = True
        plot.data_labels.number_format = number_format
        plot.data_labels.position = (
            XL_LABEL_POSITION.ABOVE if data.kind == "line" else XL_LABEL_POSITION.OUTSIDE_END
        )
        _font(plot.data_labels.font, 14)
    for index, series in enumerate(chart.series):
        color = RGBColor.from_string(PALETTE[index])
        series.format.line.color.rgb = color
        series.format.line.width = Pt(2)
        if data.kind != "line":
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color


def _draw(slide: PptxSlide, page: Page) -> None:
    entry, layout = page.entry, page.layout
    dark = layout in {"cover", "section"}
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(INK if dark else PAPER)
    if dark:
        y, top, size = _cover(entry)
        title_height = _text(slide, entry.title, MARGIN, y, BODY_WIDTH, size, WHITE, bold=True)
        if entry.subtitle:
            _text(slide, entry.subtitle, MARGIN, y + title_height + 22, BODY_WIDTH, 22, "DED8CB")
        for point in entry.points:
            top += _text(slide, point, MARGIN, top, BODY_WIDTH, 22, WHITE) + 16
        return
    top, bottom, size = _header(entry)
    title_height = _text(slide, entry.title, MARGIN, 42, BODY_WIDTH, size, bold=True)
    if entry.subtitle:
        _text(slide, entry.subtitle, MARGIN, 54 + title_height, BODY_WIDTH, 18, MUTED)
    if layout == "comparison":
        body_top = top + _column_heading(entry)
        width = (BODY_WIDTH - 52) / 2
        for index, column in enumerate(entry.columns):
            x = MARGIN + index * (width + 52)
            _text(slide, column.title, x, top, width, 24, GOLD, bold=True)
            y = body_top
            for point in column.points:
                y += _text(slide, point, x, y, width, 22, bullet=True) + 16
    elif layout in {"timeline", "process"}:
        _draw_steps(slide, page, top, bottom - top)
    elif layout == "quote":
        quote_height = _text(slide, entry.quote, MARGIN + 40, top, BODY_WIDTH - 40, 26)
        if entry.attribution:
            _text(
                slide,
                entry.attribution,
                MARGIN + 40,
                top + quote_height + 20,
                BODY_WIDTH - 40,
                18,
                MUTED,
            )
    elif layout == "table":
        _draw_table(slide, entry, top, page.column_widths)
    elif layout == "chart":
        _draw_chart(slide, entry, top, bottom - top)
    else:
        for point in entry.points:
            top += (
                _text(
                    slide,
                    point,
                    MARGIN,
                    top,
                    BODY_WIDTH,
                    24,
                    bullet=layout != "summary",
                    bold=layout == "summary",
                )
                + 16
            )


def presentation_file(spec: FileSpec) -> bytes:
    pages: list[Page] = []
    for entry in spec.slides:
        pages.extend(_pages(entry))
        if len(pages) > MAX_PAGES:
            raise AppError("presentation_too_large", "演示页数过多，请精简内容或分批生成")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Pt(WIDTH), Pt(HEIGHT)
    presentation.core_properties.title = spec.title or spec.name
    for index, page in enumerate(pages, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _draw(slide, page)
        dark = page.layout in {"cover", "section"}
        color = "DED8CB" if dark else MUTED
        if page.entry.source:
            source = "来源：" + page.entry.source
            if _height(source, BODY_WIDTH - 100, 10) > 30:
                source = "来源及版本：详见本页备注"
            _text(slide, source, MARGIN, 477, BODY_WIDTH - 100, 10, color)
        footer = f"{index:02d} / {len(pages):02d}"
        if page.parts > 1:
            footer = f"本节 {page.part}/{page.parts}   ·   " + footer
        _text(slide, footer, WIDTH - 260, 511, 208, 10, color)
        slide.shapes[-1].text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        notes = page.entry.notes
        if page.entry.source:
            notes += "\n\n来源及版本：\n" + page.entry.source
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
