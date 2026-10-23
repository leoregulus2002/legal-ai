import csv
import json
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from threading import Lock
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.common.exceptions import AppError
from app.modules.files.presentations import presentation_file
from app.modules.files.schemas import FileSpec, GenerateFilesInput

MAX_FILE_BYTES = 20 * 1024 * 1024
_pdf_font_lock = Lock()


@dataclass(frozen=True)
class RenderedFile:
    name: str
    mime_type: str
    content: bytes


def filename(value: str, extension: str) -> str:
    name = re.sub(r"[\\/:*?\"<>|\x00-\x1f\x7f]", "_", value).strip(" .")
    if name.lower().endswith("." + extension):
        name = name[: -len(extension) - 1]
    return (name[:140].strip(" .") or "生成文件") + "." + extension


def word_file(spec: FileSpec) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "List Bullet",
        "List Number",
    ):
        style = document.styles[name]
        style.font.name = "宋体"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        if name == "Normal":
            style.font.size = Pt(11)
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_after = Pt(6)
    if spec.title:
        document.add_heading(spec.title, 0)
    for block in spec.blocks:
        if block.kind == "page_break":
            document.add_page_break()
        elif block.kind == "heading":
            document.add_heading(block.text, block.level)
        elif block.kind == "table":
            columns = max(map(len, block.rows), default=0)
            if not columns:
                continue
            table = document.add_table(rows=0, cols=columns)
            table.style = "Table Grid"
            for index, values in enumerate(block.rows):
                cells = table.add_row().cells
                for column, value in enumerate(values):
                    cells[column].text = value
                    if index == 0:
                        for run in cells[column].paragraphs[0].runs:
                            run.bold = True
                if index == 0:
                    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        else:
            style_name = {"bullet": "List Bullet", "numbered": "List Number"}.get(block.kind)
            document.add_paragraph(block.text, style_name)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def spreadsheet_file(spec: FileSpec) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for entry in spec.sheets:
        sheet = workbook.create_sheet(entry.name)
        for values in entry.rows:
            sheet.append(values)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # 表格资料中的等号只是文本；只有独立声明的公式才执行。
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for address, formula in entry.formulas.items():
            sheet[address] = formula
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="294B63")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.columns, 1):
            width = max(
                (
                    sum(2 if ord(char) > 127 else 1 for char in str(cell.value or ""))
                    for cell in column
                ),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(55, max(12, width + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_file(spec: FileSpec) -> bytes:
    with _pdf_font_lock:
        if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle(
        "Body", fontName="STSong-Light", fontSize=11, leading=18, wordWrap="CJK", spaceAfter=8
    )
    heading = ParagraphStyle(
        "Heading", parent=style, fontSize=16, leading=24, spaceBefore=12, keepWithNext=True
    )
    title = ParagraphStyle("Title", parent=heading, fontSize=22, leading=30, spaceAfter=18)
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=60,
        leftMargin=60,
        topMargin=55,
        bottomMargin=55,
        title=spec.title,
    )

    def paragraph(text: str, selected: ParagraphStyle = style) -> Paragraph:
        return Paragraph(escape(text).replace("\n", "<br/>"), selected)

    story: list[object] = []
    if spec.title:
        story.append(paragraph(spec.title, title))
    number = 0
    for block in spec.blocks:
        if block.kind == "page_break":
            if story:
                story.append(PageBreak())
        elif block.kind == "table":
            count = max(map(len, block.rows), default=0)
            if count:
                rows = [
                    [paragraph(value) for value in row] + [""] * (count - len(row))
                    for row in block.rows
                ]
                table = Table(
                    rows, colWidths=[(A4[0] - 120) / count] * count, repeatRows=1, splitInRow=1
                )
                table.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2f5")),
                            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8c4ce")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ]
                    )
                )
                story.extend([table, Spacer(1, 10)])
        elif block.kind == "heading":
            story.append(paragraph(block.text, heading))
        else:
            prefix = "• " if block.kind == "bullet" else ""
            if block.kind == "numbered":
                number += 1
                prefix = f"{number}. "
            story.append(paragraph(prefix + block.text))
    document.build(story)
    return output.getvalue()


def render(spec: FileSpec) -> RenderedFile:
    binary = {
        "docx": (
            word_file,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        "pptx": (
            presentation_file,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
        "xlsx": (
            spreadsheet_file,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        "pdf": (pdf_file, "application/pdf"),
    }
    if spec.format in binary:
        renderer, mime = binary[spec.format]
        content = renderer(spec)
    elif spec.format == "csv":
        output = StringIO(newline="")
        writer = csv.writer(output)
        for row in spec.sheets[0].rows:
            writer.writerow(
                [
                    "'" + value
                    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))
                    else value
                    for value in row
                ]
            )
        content, mime = output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    else:
        text = spec.text
        if spec.format == "json":
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2, allow_nan=False)
        content, mime = text.encode("utf-8"), "text/plain; charset=utf-8"
        if spec.format == "json":
            mime = "application/json"
    if len(content) > MAX_FILE_BYTES:
        raise AppError("generated_file_too_large", "生成文件过大，请减少内容后重试")
    return RenderedFile(filename(spec.name, spec.format), mime, content)


def render_files(data: GenerateFilesInput) -> list[RenderedFile]:
    files = [render(spec) for spec in data.files]
    if len({file.name.casefold() for file in files}) != len(files):
        raise AppError("duplicate_file_name", "生成文件名称重复，请调整后重试")
    if sum(len(file.content) for file in files) > MAX_FILE_BYTES:
        raise AppError("generated_files_too_large", "生成文件总大小超限，请分批生成")
    if data.archive_name:
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            for file in files:
                archive.writestr(file.name, file.content)
        files.append(
            RenderedFile(filename(data.archive_name, "zip"), "application/zip", output.getvalue())
        )
    return files
